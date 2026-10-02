# -*- coding: utf-8 -*-
"""Document import -> persisted content -> reader position and spoken prose."""
import json
import os
import struct
import tempfile
import threading
import unittest
from pathlib import Path

from docx import Document

from novelreader import book_loader
from novelreader.document_text import word_binary_text, DocumentTextError
from novelreader.import_service import LibraryImportService
from novelreader.reader_service import ReaderService
from novelreader.storage import Storage


def word_streams():
    """Mixed Unicode/ANSI pieces at disjoint offsets, with a footer story."""
    first = "第一章 测试\r中文正文。\r"
    second = 'English text.\r\x13HYPERLINK "https://example.com"\x14Link label\x15\r'
    footer = "不应读到页脚\r"
    word = bytearray(4096)
    struct.pack_into("<HH", word, 0, 0xA5EC, 0xC1)
    struct.pack_into("<H", word, 10, 0x200)
    struct.pack_into("<H", word, 32, 14)
    struct.pack_into("<H", word, 62, 22)
    struct.pack_into("<I", word, 76, len(first) + len(second))
    struct.pack_into("<H", word, 152, 93)
    word[2048:2048 + len(first) * 2] = first.encode("utf-16-le")
    word[1024:1024 + len(second)] = second.encode("cp1252")
    word[3072:3072 + len(footer) * 2] = footer.encode("utf-16-le")
    positions = [0, len(first), len(first) + len(second), len(first) + len(second) + len(footer)]
    plc = bytearray(struct.pack("<4I", *positions))
    for offset in (2048, 0x40000000 | 2048, 3072):
        plc.extend(struct.pack("<HIH", 0, offset, 0))
    clx = b"\x02" + struct.pack("<I", len(plc)) + plc
    struct.pack_into("<II", word, 418, 0, len(clx))
    return bytes(word), clx.ljust(4096, b"\x00")


def write_doc(path, word, table):
    """Minimal real CFB container with two regular 4096-byte streams."""
    end, free = 0xFFFFFFFE, 0xFFFFFFFF
    header = bytearray(512)
    header[:8] = bytes.fromhex("D0CF11E0A1B11AE1")
    struct.pack_into("<HHHH", header, 24, 0x3E, 3, 0xFFFE, 9)
    struct.pack_into("<H", header, 32, 6)
    struct.pack_into("<II", header, 44, 1, 0)
    struct.pack_into("<I", header, 56, 4096)
    struct.pack_into("<IIII", header, 60, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, 17, *([free] * 108))
    directory = bytearray(512)
    for index, (name, kind, start, size) in enumerate([
        ("Root Entry", 5, end, 0), ("WordDocument", 2, 1, 4096), ("1Table", 2, 9, 4096),
    ]):
        offset = index * 128
        encoded = (name + "\x00").encode("utf-16-le")
        directory[offset:offset + len(encoded)] = encoded
        struct.pack_into("<HBBIII", directory, offset + 64, len(encoded), kind, 1,
                         free, 2 if index == 1 else free, 1 if index == 0 else free)
        struct.pack_into("<IQ", directory, offset + 116, start, size)
    fat = [end] + list(range(2, 9)) + [end] + list(range(10, 17)) + [end, 0xFFFFFFFD]
    fat.extend([free] * (128 - len(fat)))
    path.write_bytes(header + directory + word + table + struct.pack("<128I", *fat))


class DocumentReadingTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.library_path = self.root / "data" / "library.json"
        from unittest.mock import patch
        environment = patch.dict(os.environ, {"DOUBAO_NOVEL_DATA": os.fspath(self.library_path.parent)})
        environment.start()
        self.addCleanup(environment.stop)
        self.service = LibraryImportService(self.library_path)

    def import_and_read(self, source):
        candidates = self.service.inspect_files([os.fspath(source)])
        self.assertTrue(candidates[0].supported)
        result = self.service.import_files(candidates, "overwrite", True, threading.Event(), lambda _: None)
        self.assertEqual(result["succeeded"], 1, result)
        reader = ReaderService(self.library_path)
        opened = reader.open_book(result["openAfterImportBookId"])
        return reader, opened, result["openAfterImportBookId"]

    def test_markdown_spoken_prose_and_heading_chapters(self):
        source = self.root / "article.md"
        source.write_text('''---
author: 不朗读的元数据
---
# 阅读标题
普通**粗体**与*斜体*、~~删除线~~。[链接文字](https://example.com/path)
![图片说明](image.png)
> 引用文字。
- 列表文字。
- [x] 完成项目。
```python
不朗读的代码
```

    不朗读的缩进代码

## 第二节
第二节正文 <strong>加粗文字</strong>。
内嵌<script>不朗读的脚本</script>文字。
| 表头 | 内容 |
| --- | --- |
| 项目 | 表格正文 |
''', encoding="utf-8")
        book = book_loader.parse_book(os.fspath(source))
        self.assertEqual(book.title, "阅读标题")
        self.assertEqual([c.title for c in book.chapters], ["阅读标题", "第二节"])
        spoken = "\n".join(c.tts_content()[0] for c in book.chapters)
        for expected in ("普通粗体与斜体", "链接文字", "引用文字", "列表文字", "完成项目", "加粗文字", "表格正文"):
            self.assertIn(expected, spoken)
        for excluded in ("https://", "图片说明", "不朗读", "```", "**", "[x]", "<strong>", "#"):
            self.assertNotIn(excluded, spoken)
        reader, opened, book_id = self.import_and_read(source)
        reader.update_position(opened["sessionId"], 1, 2)
        reopened = ReaderService(self.library_path).open_book(book_id)
        self.assertEqual(reopened["position"]["chapterIndex"], 1)
        self.assertEqual(reopened["position"]["charOffset"], 2)

    def test_markdown_only_code_is_not_imported(self):
        source = self.root / "code.markdown"
        source.write_text("```python\nprint('code')\n```", encoding="utf-8")
        result = self.service.import_files(self.service.inspect_files([os.fspath(source)]), "overwrite", True,
                                         threading.Event(), lambda _: None)
        self.assertEqual(result["failed"], 1)
        self.assertIn("未提取到", result["results"][0]["error"]["message"])

    def test_doc_reads_piece_order_chinese_fields_and_excludes_footer(self):
        word, table = word_streams()
        source = self.root / "legacy.doc"
        write_doc(source, word, table)
        book = book_loader.parse_book(os.fspath(source))
        self.assertEqual(book.format, "doc")
        text = "\n".join(c.content for c in book.chapters)
        self.assertIn("中文正文。\nEnglish text.\nLink label", text)
        self.assertNotIn("HYPERLINK", text)
        self.assertNotIn("https://", text)
        self.assertNotIn("页脚", text)
        self.import_and_read(source)

    def test_doc_rejects_encryption_and_out_of_bounds_piece(self):
        word, table = word_streams()
        encrypted = bytearray(word)
        struct.pack_into("<H", encrypted, 10, 0x300)
        source = self.root / "encrypted.doc"
        write_doc(source, encrypted, table)
        result = self.service.import_files(self.service.inspect_files([os.fspath(source)]), "overwrite", True,
                                         threading.Event(), lambda _: None)
        self.assertEqual(result["failed"], 1)
        self.assertIn("加密", result["results"][0]["error"]["message"])
        broken = bytearray(table)
        struct.pack_into("<I", broken, 5 + 16 + 2, 100000)
        with self.assertRaises(DocumentTextError):
            word_binary_text(word, broken)

    def test_docx_tables_are_read_in_document_order(self):
        source = self.root / "table.docx"
        document = Document()
        document.add_paragraph("表格之前。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "第一列内容。"
        table.cell(0, 1).text = "第二列内容。"
        document.add_paragraph("表格之后。")
        document.save(source)
        book = book_loader.parse_book(os.fspath(source))
        text = "\n".join(c.content for c in book.chapters)
        self.assertEqual(text, "表格之前。\n第一列内容。\n第二列内容。\n表格之后。")
        self.import_and_read(source)

    def test_unexpected_parser_error_remains_path_free(self):
        from unittest.mock import patch
        source = self.root / "private.docx"
        source.write_text("damaged")
        with patch("novelreader.book_loader.parse_book", side_effect=ValueError(os.fspath(source))):
            result = self.service.import_files(self.service.inspect_files([os.fspath(source)]), "overwrite", True,
                                             threading.Event(), lambda _: None)
        self.assertNotIn(os.fspath(source), json.dumps(result))


if __name__ == "__main__":
    unittest.main()
